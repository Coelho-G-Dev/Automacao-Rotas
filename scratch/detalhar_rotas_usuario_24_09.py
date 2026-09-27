import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import openpyxl
import pandas as pd

wb = openpyxl.load_workbook(r"C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 24-09.xlsx", data_only=True)

# 1. Carregar PESO_CAP
ws_pcap = wb["PESO_CAP"]
lojas_pcap = {}
for r in range(6, ws_pcap.max_row + 1):
    loja = ws_pcap.cell(r, 2).value
    peso = ws_pcap.cell(r, 5).value
    rota = ws_pcap.cell(r, 6).value
    ger = ws_pcap.cell(r, 8).value
    if loja and peso is not None:
        cod = str(loja).split(" - ")[0].strip()
        lojas_pcap[cod] = {
            "nome": str(loja),
            "peso": float(peso),
            "rota_definida": str(rota),
            "geracao": str(ger)
        }

print("=== LOJAS DO PESO_CAP (USUÁRIO) ===")
for cod, d in lojas_pcap.items():
    print(f"Loja {cod:>4} | Rota: {d['rota_definida']:<5} | Geração: {d['geracao']:<5} | Peso: {d['peso']:>9.2f} kg | Nome: {d['nome']}")

# 2. Carregar ROTAS_CAP
ws_rcap = wb["ROTAS_CAP"]
print("\n=== VEÍCULOS ROTAS_CAP (USUÁRIO) ===")
for r in range(8, ws_rcap.max_row + 1):
    vals = [ws_rcap.cell(r, c).value for c in range(2, 11)]
    paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in vals[:4] if p is not None]
    veic = vals[6]
    obs = vals[7]
    if paradas:
        p_calc = sum(lojas_pcap.get(p, {}).get("peso", 0.0) for p in paradas if p != "96")
        if "96" in paradas:
            # 96 aparece em 2 carros
            p_calc += lojas_pcap.get("96", {}).get("peso", 0.0) / 2
        print(f"Linha {r:02d} | Veículo: {str(veic):<7} | Peso Calculado: {p_calc:>9.2f} kg | Paradas: {paradas} | Obs: {obs}")

# 3. Carregar ROTA_INT
ws_rint = wb["ROTA_INT"]
print("\n=== VEÍCULOS ROTA_INT (USUÁRIO) ===")
for r in range(8, ws_rint.max_row + 1):
    vals = [ws_rint.cell(r, c).value for c in range(2, 11)]
    paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in vals[:4] if p is not None]
    veic = vals[6]
    obs = vals[5]
    if paradas:
        print(f"Linha {r:02d} | Veículo: {str(veic):<7} | Paradas: {paradas} | Obs: {obs}")

# 4. CROSSDOCK
ws_cd = wb["CROSSDOCK"]
print("\n=== VEÍCULOS CROSSDOCK (USUÁRIO) ===")
for r in range(5, ws_cd.max_row + 1):
    vals = [ws_cd.cell(r, c).value for c in range(2, 12)]
    paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in vals[:7] if p is not None]
    veic = vals[7]
    if paradas:
        print(f"Linha {r:02d} | Veículo: {str(veic):<7} | Paradas: {paradas}")
