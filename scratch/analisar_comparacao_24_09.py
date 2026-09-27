import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import openpyxl
import pandas as pd

wb_user = openpyxl.load_workbook(r"C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 24-09.xlsx", data_only=True)
wb_auto = openpyxl.load_workbook(r"C:\Users\gabri\Desktop\Separacao\Prototipo_Rotas_Frios_Novo.xlsx", data_only=True)

# 1. Carregar lojas e pesos no PESO_CAP do usuario
ws_pcap = wb_user["PESO_CAP"]
user_lojas = {}
for r in range(6, ws_pcap.max_row + 1):
    loja = ws_pcap.cell(r, 2).value
    peso = ws_pcap.cell(r, 5).value
    rota_col = ws_pcap.cell(r, 6).value
    geracao = ws_pcap.cell(r, 8).value
    if loja and peso is not None:
        cod = str(loja).split(" - ")[0].strip()
        user_lojas[cod] = {
            "nome": str(loja),
            "peso": float(peso),
            "rota_col": str(rota_col),
            "geracao": str(geracao)
        }

# 2. Carregar rotas manuais do usuario em ROTAS_CAP
ws_rcap = wb_user["ROTAS_CAP"]
user_rotas = []
for r in range(8, ws_rcap.max_row + 1):
    p1 = ws_rcap.cell(r, 2).value
    p2 = ws_rcap.cell(r, 3).value
    p3 = ws_rcap.cell(r, 4).value
    p4 = ws_rcap.cell(r, 5).value
    veic = ws_rcap.cell(r, 8).value
    paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in [p1, p2, p3, p4] if p is not None]
    if paradas:
        user_rotas.append({
            "veiculo": str(veic),
            "paradas": paradas,
            "linha": r
        })

# 3. Carregar rotas da automacao em ROTAS_CAP
ws_acap = wb_auto["ROTAS_CAP"]
auto_rotas = []
for r in range(8, ws_acap.max_row + 1):
    p1 = ws_acap.cell(r, 2).value
    p2 = ws_acap.cell(r, 3).value
    p3 = ws_acap.cell(r, 4).value
    p4 = ws_acap.cell(r, 5).value
    veic = ws_acap.cell(r, 7).value
    peso = ws_acap.cell(r, 8).value
    rota_padrao = ws_acap.cell(r, 9).value
    obs = ws_acap.cell(r, 10).value
    paradas = [str(p).split()[0].strip() for p in [p1, p2, p3, p4] if p is not None and str(p).strip() != ""]
    if paradas:
        auto_rotas.append({
            "veiculo": str(veic),
            "peso": float(peso or 0),
            "rota_padrao": str(rota_padrao),
            "paradas": paradas,
            "obs": str(obs)
        })

print("=" * 80)
print("ROTAS MANUAIS FEITAS HOJE PELO USUÁRIO (24/09):")
print("=" * 80)
total_peso_user = 0.0
for idx, r in enumerate(user_rotas, 1):
    # Calcular peso
    p_acum = 0.0
    for p in r["paradas"]:
        if p in user_lojas:
            p_acum += user_lojas[p]["peso"]
        elif p == "96":
            p_acum += user_lojas.get("96", {}).get("peso", 0.0) / 2 # split
    total_peso_user += p_acum
    print(f"U{idx:02d} | Veículo: {r['veiculo']:<7} | Peso Est: {p_acum:>9.2f} kg | Paradas: {r['paradas']}")

print(f"\nTotal Carros Usuário: {len(user_rotas)}")

print("\n" + "=" * 80)
print("ROTAS GERADAS PELO ALGORITMO:")
print("=" * 80)
for idx, r in enumerate(auto_rotas, 1):
    print(f"A{idx:02d} | Veículo: {r['veiculo']:<7} | Peso: {r['peso']:>10.2f} kg | Rota Padrão: {r['rota_padrao']:<7} | Paradas: {r['paradas']}")
print(f"\nTotal Carros Algoritmo: {len(auto_rotas)}")
