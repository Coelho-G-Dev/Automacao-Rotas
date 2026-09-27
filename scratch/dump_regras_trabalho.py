import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

path = r'C:\Users\gabri\Desktop\Trabalho\Work\ROTAS DO INTERIOR.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    print("=" * 100)
    print(f"ABA COMPLETA: {sheet}")
    print("=" * 100)
    for r in range(1, ws.max_row + 1):
        row_vals = []
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val is not None:
                row_vals.append(f"C{c}: {str(val).strip()}")
        if row_vals:
            print(f"Linha {r:02d} | " + " | ".join(row_vals))
