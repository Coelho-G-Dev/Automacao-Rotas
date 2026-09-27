import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

path = r'C:\Users\gabri\Desktop\Trabalho\Work\ROTAS DO INTERIOR.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
print("Aba(s):", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n=== ABA: {sname} (max_row={ws.max_row}, max_col={ws.max_column}) ===")
    for r in range(1, min(60, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"L{r:02d}: {[v for v in vals if v is not None]}")
