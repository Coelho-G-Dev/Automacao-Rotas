import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

path = r'C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 26-09  - .xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

print("=" * 80)
print("ROTAS MANUAIS DO INTERIOR (26/09):")
print("=" * 80)
ws_rint = wb['ROTA_INT']
for r in range(1, ws_rint.max_row + 1):
    vals = [ws_rint.cell(r, c).value for c in range(1, 15)]
    if any(v is not None for v in vals):
        print(f"Linha {r:02d}: {[v for v in vals if v is not None]}")

print("\n" + "=" * 80)
print("LOJAS PESO__INT (26/09):")
print("=" * 80)
ws_pint = wb['PESO__INT']
for r in range(5, ws_pint.max_row + 1):
    loja = ws_pint.cell(r, 2).value
    peso = ws_pint.cell(r, 5).value
    rota = ws_pint.cell(r, 6).value
    ger = ws_pint.cell(r, 8).value
    if loja is not None and peso is not None:
        print(f"Loja: {str(loja):<50} | Peso: {float(peso):>10.2f} kg | Rota: {str(rota):<5} | Ger: {str(ger)}")

print("\n" + "=" * 80)
print("CROSSDOCK (26/09):")
print("=" * 80)
ws_cd = wb['CROSSDOCK']
for r in range(1, ws_cd.max_row + 1):
    vals = [ws_cd.cell(r, c).value for c in range(1, 15)]
    if any(v is not None for v in vals):
        print(f"Linha {r:02d}: {[v for v in vals if v is not None]}")
