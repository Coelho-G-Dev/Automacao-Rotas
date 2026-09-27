import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import openpyxl
import pandas as pd
from automacao import obter_base_enderecos, carregar_dados_frios, localizar_aba_e_cabecalho

aba, row = localizar_aba_e_cabecalho('Acomp_EvoluSep.xlsx')
df = carregar_dados_frios('Acomp_EvoluSep.xlsx', aba, row)
base = obter_base_enderecos(df, 'Cadastro_Lojas_Enderecos.xlsx')

wb_user = openpyxl.load_workbook(r"C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 24-09.xlsx", data_only=True)
ws_pcap = wb_user["PESO_CAP"]
lojas_pcap = {}
for r in range(6, ws_pcap.max_row + 1):
    loja = ws_pcap.cell(r, 2).value
    peso = ws_pcap.cell(r, 5).value
    rota = ws_pcap.cell(r, 6).value
    ger = ws_pcap.cell(r, 8).value
    if loja and peso is not None:
        cod = str(loja).split(" - ")[0].strip()
        lojas_pcap[cod] = {"peso": float(peso), "rota": str(rota), "ger": str(ger), "nome": str(loja)}

ws_rcap = wb_user["ROTAS_CAP"]
print("=" * 90)
print("ANÁLISE DETALHADA DAS 12 ROTAS MANUAIS DA CAPITAL FEITAS HOJE PELO USUÁRIO (24/09)")
print("=" * 90)

faixas = {
    "3/4": (5500, 6000),
    "3/4.": (5500, 6000),
    "TOCO": (8000, 9000),
    "TRUCK": (12000, 14000),
    "BITRUCK": (15500, 16000),
    "CARRETA": (24000, 28000)
}

for r in range(8, ws_rcap.max_row + 1):
    vals = [ws_rcap.cell(r, c).value for c in range(2, 11)]
    paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in vals[:4] if p is not None]
    veic = str(vals[6]).strip().upper()
    if paradas:
        detalhes = []
        peso_carro = 0.0
        rotas_envolvidas = set()
        for p in paradas:
            info_cad = base.get(p, {})
            bairro = info_cad.get("Bairro", "Bairro Desconhecido")
            p_loja = lojas_pcap.get(p, {}).get("peso", 0.0)
            rota_p = lojas_pcap.get(p, {}).get("rota", "N/D")
            ger = lojas_pcap.get(p, {}).get("ger", "N/D")
            rotas_envolvidas.add(rota_p)
            if p == "96":
                # verificar split
                p_loja = 9442.87 if r == 15 else (18885.73 - 9442.87)
            peso_carro += p_loja
            detalhes.append(f"L{p} [{bairro} | {p_loja:,.1f}kg | Rota {rota_p} | {ger}]")
        
        f_min, f_max = faixas.get(veic, (0, 99999))
        status_peso = "DENTRO DA FAIXA" if f_min <= peso_carro <= f_max else ("ABAIXO DO PISO" if peso_carro < f_min else "ACIMA DO TETO")
        dif_piso = peso_carro - f_min
        
        print(f"\n[CARRO {r-7:02d}] Veículo: {veic:<7} | Peso: {peso_carro:>10.2f} kg | Status: {status_peso} (Faixa {f_min:,.0f} a {f_max:,.0f} kg, Dif: {dif_piso:+,.1f} kg)")
        print(f"   Rotas PESO_CAP: {list(rotas_envolvidas)}")
        for d in detalhes:
            print(f"      • {d}")

print("\n" + "=" * 90)
