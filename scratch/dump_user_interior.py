import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook(r'C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 24-09.xlsx', data_only=True)

print("=" * 80)
print("ABA: ROTA_INT (Todas as linhas preenchidas)")
print("=" * 80)
ws_rint = wb['ROTA_INT']
for r in range(1, ws_rint.max_row + 1):
    vals = [ws_rint.cell(r, c).value for c in range(1, ws_rint.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Linha {r:02d}: {[v for v in vals if v is not None]}")

print("\n" + "=" * 80)
print("ABA: PESO__INT (Todas as lojas cadastradas no interior pelo usuário)")
print("=" * 80)
ws_pint = wb['PESO__INT']
for r in range(1, ws_pint.max_row + 1):
    vals = [ws_pint.cell(r, c).value for c in range(1, 12)]
    if any(v is not None for v in vals):
        print(f"Linha {r:02d}: {vals}")

print("\n" + "=" * 80)
print("ABA: CROSSDOCK")
print("=" * 80)
ws_cd = wb['CROSSDOCK']
for r in range(1, ws_cd.max_row + 1):
    vals = [ws_cd.cell(r, c).value for c in range(1, ws_cd.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"Linha {r:02d}: {[v for v in vals if v is not None]}")
