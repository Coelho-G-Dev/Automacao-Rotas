import os
import sys
import math
import warnings
import zipfile
from collections import defaultdict, Counter
import itertools
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

ARQUIVO_ENTRADA_PADRAO = r"C:\Users\gabri\Desktop\Separacao\Acomp_EvoluSep.xlsx"
ARQUIVO_SAIDA_PADRAO = r"C:\Users\gabri\Desktop\Separacao\Prototipo_Rotas_Frios.xlsx"
ARQUIVO_ENDERECOS_PADRAO = r"C:\Users\gabri\Desktop\Separacao\Cadastro_Lojas_Enderecos.xlsx"

IGNORAR_LOJAS_4_DIGITOS = True
LIMITE_MAX_LOJAS_POR_ROTA = 4
LIMITE_LOJAS_CARRETA_CAPITAL = 1
LIMITE_LOJAS_CARRETA_INTERIOR = 4
PESO_MINIMO_ROTA = 5000
TOLERANCIA_PISO_CAPITAL = 1000.0

PARAMETROS_FROTA = [
    {"tipo": "3/4",     "peso_minimo": 5500,  "piso_operacional": 5000, "capacidade": 6000,  "regioes": ["CAPITAL"]},
    {"tipo": "Toco",    "peso_minimo": 8000,  "piso_operacional": 8000, "capacidade": 9000,  "regioes": ["CAPITAL", "INTERIOR"]},
    {"tipo": "Truck",   "peso_minimo": 12000, "piso_operacional": 12000, "capacidade": 14000, "regioes": ["CAPITAL", "INTERIOR"]},
    {"tipo": "Bitruck", "peso_minimo": 15500, "piso_operacional": 15500, "capacidade": 16000, "regioes": ["CAPITAL", "INTERIOR"]},
    {"tipo": "Carreta", "peso_minimo": 24000, "piso_operacional": 24000, "capacidade": 28000, "regioes": ["CAPITAL", "INTERIOR"]}
]

CAPACIDADE_CARRETA = 28000 

SETORES_ALVO = ["CONGELADOS", "RESFRIADOS"]
ROTAS_INVALIDAS = {"", "NAN", "NONE", "NULL", "NAO CONFIGURADO", "NÃO CONFIGURADO", "SEM ROTA"}

CD_LAT_DEFAULT = -2.565
CD_LON_DEFAULT = -44.223

LOJAS_TOCO_CIMA = {"451", "454", "457", "459", "435", "434", "461"}

def formatar_codigo_filial(val) -> str:
    if pd.isna(val):
        return ""
    try:
        f = float(val)
        if f.is_integer():
            return str(int(f))
        return str(f)
    except (ValueError, TypeError):
        return str(val).strip()

def normalizar_nome_rota(rota_raw) -> str:
    if pd.isna(rota_raw):
        return ""
    
    rota_str = str(rota_raw).strip()
    
    if "ROTA" in rota_str.upper() and "/" in rota_str:
        parte_rota = rota_str.split("/")[0].strip()
        return normalizar_nome_rota(parte_rota)
        
    if rota_str.upper() == "ALL" or "ROTA ALL" in rota_str.upper():
        return "Rota All"
        
    try:
        f = float(rota_raw)
        if f.is_integer():
            rota_raw = int(f)
    except (ValueError, TypeError):
        pass
        
    rota_str = str(rota_raw).strip()
    if not rota_str.lower().startswith("rota"):
        return f"Rota {rota_str}"
    return rota_str

def extrair_setor_numero(rota_str: str) -> str:
    s = str(rota_str or "").upper().replace("ROTA", "").strip()
    digits = ''.join(filter(str.isdigit, s))
    if digits:
        return digits
    if "ALL" in s:
        return "ALL"
    return s

def calcular_distancia_cd(lat, lon) -> float:
    try:
        lat = float(lat)
        lon = float(lon)
        return math.sqrt((lat - CD_LAT_DEFAULT)**2 + (lon - CD_LON_DEFAULT)**2) * 111.0
    except (ValueError, TypeError):
        return 999.0

def obter_restricoes_loja(filial_cod, base_enderecos: dict = None) -> dict:
    f_str = str(filial_cod).strip()
    if " " in f_str:
        f_str = f_str.split()[0]
        
    try:
        f_int = int(float(f_str))
    except (ValueError, TypeError):
        f_int = None
        
    restr_cad = None
    uf_cad = ""
    nome_cad = ""
    if base_enderecos and f_str in base_enderecos:
        info = base_enderecos[f_str]
        restr_cad = info.get("Restrição Veículo")
        uf_cad = str(info.get("UF", "")).strip().upper()
        nome_cad = str(info.get("Nome Loja", "")).strip().upper()

    eh_ceara = (
        uf_cad == "CE"
        or (restr_cad and "SOMENTE CARRETA" in restr_cad.upper())
        or any(k in nome_cad for k in [" - CE", "-CE", " (CE)", " / CE", "/CE", " CE ", "CEARÁ", "CEARA"])
        or nome_cad.endswith(" CE")
        or (restr_cad and "CE" in restr_cad.upper())
    )
    if eh_ceara:
        return {
            "veiculos_permitidos": ["Carreta"],
            "veiculos_proibidos": ["3/4", "Toco", "Truck", "Bitruck"],
            "capacidade_maxima": CAPACIDADE_CARRETA,
            "descricao": "Somente Carreta"
        }

    if f_int in (411, 418) or (restr_cad and "SOMENTE 3/4" in restr_cad.upper()):
        return {
            "veiculos_permitidos": ["3/4"],
            "veiculos_proibidos": ["Toco", "Truck", "Bitruck", "Carreta"],
            "capacidade_maxima": 6000,
            "descricao": "Somente 3/4"
        }

    if f_int == 560 or (restr_cad and "SOMENTE TOCO" in restr_cad.upper()):
        return {
            "veiculos_permitidos": ["3/4", "Toco"],
            "veiculos_proibidos": ["Truck", "Bitruck", "Carreta"],
            "capacidade_maxima": 10500,
            "descricao": "Somente Toco"
        }

    if f_str in LOJAS_TOCO_CIMA or (f_int is not None and f_int in {451, 454, 457, 459, 435, 434, 461}) or (restr_cad and "TOCO PARA CIMA" in restr_cad.upper()):
        return {
            "veiculos_permitidos": ["Toco", "Truck", "Bitruck", "Carreta"],
            "veiculos_proibidos": ["3/4"],
            "capacidade_maxima": CAPACIDADE_CARRETA,
            "descricao": "De Toco para cima"
        }

    if (f_int is not None and ((400 <= f_int < 500) or f_int in (27, 415))) or (restr_cad and "NÃO ENTRA BITRUCK" in restr_cad.upper().replace("NAO", "NÃO")):
        return {
            "veiculos_permitidos": ["3/4", "Toco", "Truck", "Carreta"],
            "veiculos_proibidos": ["Bitruck"],
            "capacidade_maxima": CAPACIDADE_CARRETA,
            "descricao": "Não entra Bitruck"
        }

    return {
        "veiculos_permitidos": ["3/4", "Toco", "Truck", "Bitruck", "Carreta"],
        "veiculos_proibidos": [],
        "capacidade_maxima": CAPACIDADE_CARRETA,
        "descricao": "Sem restrição"
    }

def obter_veiculos_compativeis(filiais: list, base_enderecos: dict = None, regiao: str = None) -> list:
    if not filiais:
        return [v["tipo"] for v in PARAMETROS_FROTA]
        
    tipos_permitidos = set([v["tipo"] for v in PARAMETROS_FROTA])
    for f in filiais:
        restr = obter_restricoes_loja(f, base_enderecos)
        tipos_permitidos = tipos_permitidos.intersection(set(restr["veiculos_permitidos"]))
        
    lojas_unicas = set(str(f).split()[0].replace("(Parcial)", "").strip() for f in filiais)
    lojas_validas = [l for l in lojas_unicas if l]
    if regiao and str(regiao).upper() == "CAPITAL":
        if len(lojas_validas) > LIMITE_LOJAS_CARRETA_CAPITAL:
            tipos_permitidos.discard("Carreta")
    elif regiao and str(regiao).upper() == "INTERIOR":
        if len(lojas_validas) > LIMITE_LOJAS_CARRETA_INTERIOR:
            tipos_permitidos.discard("Carreta")
            
    return [v["tipo"] for v in PARAMETROS_FROTA if v["tipo"] in tipos_permitidos]

def obter_capacidade_maxima_conjunto(filiais: list, base_enderecos: dict = None, regiao: str = None) -> float:
    veiculos_validos = obter_veiculos_compativeis(filiais, base_enderecos, regiao)
    if not veiculos_validos:
        return 0.0
    caps = [
        v["capacidade"] for v in PARAMETROS_FROTA
        if v["tipo"] in veiculos_validos and (not regiao or regiao.upper() in v.get("regioes", ["CAPITAL", "INTERIOR"]))
    ]
    return max(caps) if caps else CAPACIDADE_CARRETA

def selecionar_menor_veiculo(
    peso: float,
    filiais: list = None,
    regiao: str = "CAPITAL",
    base_enderecos: dict = None,
    permitir_excecao: bool = True
):
    veiculos_validos = obter_veiculos_compativeis(filiais or [], base_enderecos, regiao)
    if not veiculos_validos:
        veiculos_validos = [v["tipo"] for v in PARAMETROS_FROTA]
        
    lojas_unicas = set(str(f).split()[0].replace("(Parcial)", "").strip() for f in (filiais or []))
    lojas_validas = [l for l in lojas_unicas if l]
    if regiao and str(regiao).upper() == "CAPITAL":
        if len(lojas_validas) > LIMITE_LOJAS_CARRETA_CAPITAL and "Carreta" in veiculos_validos:
            veiculos_validos = [v for v in veiculos_validos if v != "Carreta"]
    elif regiao and str(regiao).upper() == "INTERIOR":
        if len(lojas_validas) > LIMITE_LOJAS_CARRETA_INTERIOR and "Carreta" in veiculos_validos:
            veiculos_validos = [v for v in veiculos_validos if v != "Carreta"]

    for veiculo in PARAMETROS_FROTA:
        if veiculo["tipo"] not in veiculos_validos:
            continue
        if regiao and regiao.upper() not in veiculo.get("regioes", ["CAPITAL", "INTERIOR"]):
            continue
        
        piso = veiculo.get("piso_operacional", veiculo["peso_minimo"]) if (regiao and regiao.upper() == "CAPITAL") else veiculo["peso_minimo"]
        if piso <= peso <= veiculo["capacidade"]:
            return veiculo["tipo"], False, 0.0

    if regiao and str(regiao).upper() == "CAPITAL" and permitir_excecao:
        for veiculo in PARAMETROS_FROTA:
            if veiculo["tipo"] not in veiculos_validos:
                continue
            if "CAPITAL" not in veiculo.get("regioes", ["CAPITAL", "INTERIOR"]):
                continue
            
            piso_ref = veiculo["peso_minimo"]
            piso_tol = piso_ref - TOLERANCIA_PISO_CAPITAL
            
            if piso_tol <= peso <= veiculo["capacidade"]:
                falta = max(0.0, piso_ref - peso)
                return veiculo["tipo"], True, falta

    return None, False, 0.0

def obter_menor_veiculo_alvo(filiais: list = None, base_enderecos: dict = None, regiao: str = "CAPITAL") -> dict:
    veiculos_validos = obter_veiculos_compativeis(filiais or [], base_enderecos, regiao)
    for v in PARAMETROS_FROTA:
        if v["tipo"] in veiculos_validos and (not regiao or regiao.upper() in v.get("regioes", ["CAPITAL", "INTERIOR"])):
            return v
    return PARAMETROS_FROTA[0]

def ordenar_paradas(filiais: list, regiao: str, base_enderecos: dict) -> list:
    def chave_ordenacao(f):
        f_limpo = str(f).split()[0].strip()

        if f_limpo in ["60", "61", "90"]:
            return (0, 0)

        if f_limpo.startswith("14"):
            return (2, calcular_distancia_cd(
                base_enderecos.get(f_limpo, {}).get("Latitude"),
                base_enderecos.get(f_limpo, {}).get("Longitude")
            ))

        info = base_enderecos.get(f_limpo, {})
        d = calcular_distancia_cd(info.get("Latitude"), info.get("Longitude"))
        return (1, d)
        
    return sorted(filiais, key=chave_ordenacao)

def encontrar_coluna(colunas, prioridades):
    colunas_lista = [str(c).strip() for c in colunas]
    colunas_upper = {c.upper(): c for c in colunas_lista}
    
    for termo in prioridades:
        t_upper = termo.upper()
        if t_upper in colunas_upper:
            return colunas_upper[t_upper]
            
    for termo in prioridades:
        t_upper = termo.upper()
        for col_upper, col_orig in colunas_upper.items():
            if col_upper.endswith(f".{t_upper}") or col_upper == t_upper:
                return col_orig

    for termo in prioridades:
        t_upper = termo.upper()
        for col_upper, col_orig in colunas_upper.items():
            if t_upper in col_upper:
                return col_orig
                
    return None

def inferir_dados_localizacao(cliente_str: str, regiao: str):
    c = str(cliente_str).upper()
    cidade = "São Luís" if regiao == "CAPITAL" else "Interior"
    uf = "MA"
    bairro = ""
    
    if any(k in c for k in [
        "MARACANAU", "HENRIQUE JORGE", "MARANGUAPE", "CANINDE", "ARACATI", 
        "ZE WALTER", "CAUCAIA", "MESSEJANA", "TIANGUA", "SOBRAL", "ITAPIPOCA", 
        "CRATEUS", "QUIXERAMOBIM", "JUAZEIRO", "RUSSAS", "FORTALEZA",
        " - CE", "-CE", " (CE)", " / CE", "/CE", " CE ", "CEARÁ", "CEARA"
    ]) or c.endswith(" CE"):
        uf = "CE"
        if "MARACANAU" in c: cidade = "Maracanaú"; bairro = "Centro"
        elif "HENRIQUE JORGE" in c: cidade = "Fortaleza"; bairro = "Henrique Jorge"
        elif "MARANGUAPE" in c: cidade = "Maranguape"; bairro = "Centro"
        elif "CANINDE" in c: cidade = "Canindé"; bairro = "Centro"
        elif "ARACATI" in c: cidade = "Aracati"; bairro = "Centro"
        elif "ZE WALTER" in c: cidade = "Fortaleza"; bairro = "Prefeito José Walter"
        elif "CAUCAIA" in c: cidade = "Caucaia"; bairro = "Centro"
        elif "MESSEJANA" in c: cidade = "Fortaleza"; bairro = "Messejana"
        elif "TIANGUA" in c: cidade = "Tianguá"; bairro = "Centro"
        elif "SOBRAL" in c: cidade = "Sobral"; bairro = "Centro"
        elif "ITAPIPOCA" in c: cidade = "Itapipoca"; bairro = "Centro"
        elif "CRATEUS" in c: cidade = "Crateús"; bairro = "Centro"
        elif "QUIXERAMOBIM" in c: cidade = "Quixeramobim"; bairro = "Centro"
        elif "JUAZEIRO" in c: cidade = "Juazeiro do Norte"; bairro = "Centro"
        elif "RUSSAS" in c: cidade = "Russas"; bairro = "Centro"
        elif "FORTALEZA" in c: cidade = "Fortaleza"; bairro = "Centro"
        else: cidade = "Fortaleza"; bairro = "Centro"
    elif any(k in c for k in ["PIRIPIRI", "TERESINA", "TIMON"]):
        if "PIRIPIRI" in c: cidade = "Piripiri"; uf = "PI"; bairro = "Centro"
        elif "TERESINA" in c: cidade = "Teresina"; uf = "PI"; bairro = "Uruguai / Novafapi"
        elif "TIMON" in c: cidade = "Timon"; uf = "MA"; bairro = "Alvorada"
    else:

        if "CHAPADINHA" in c: cidade = "Chapadinha"; bairro = "Centro"
        elif "PINHEIRO" in c: cidade = "Pinheiro"; bairro = "Centro"
        elif "CODO" in c: cidade = "Codó"; bairro = "Centro"
        elif "SAO MATEUS" in c: cidade = "São Mateus do Maranhão"; bairro = "Centro"
        elif "ROSARIO" in c: cidade = "Rosário"; bairro = "Centro"
        elif "CAXIAS" in c: cidade = "Caxias"; bairro = "Centro"
        elif "ARARI" in c: cidade = "Arari"; bairro = "Centro"
        elif "VIANA" in c: cidade = "Viana"; bairro = "Centro"
        elif "BARREIRINHAS" in c: cidade = "Barreirinhas"; bairro = "Centro"
        elif "TUTOIA" in c: cidade = "Tutóia"; bairro = "Centro"
        elif "SAO BENTO" in c: cidade = "São Bento"; bairro = "Centro"
        elif "CURURUPU" in c: cidade = "Cururupu"; bairro = "Centro"
        elif "CASA DO ARROZ" in c: cidade = "Santa Luzia"; bairro = "Centro"

        elif "COHAB" in c and "IV" not in c: bairro = "Cohab Anil"
        elif "COHAMA" in c: bairro = "Cohama"
        elif "TURU-SM08" in c or "TURU" in c: bairro = "Turu"
        elif "RIO ANIL" in c: bairro = "Vila Palmeira / Rio Anil"
        elif "CALHAU" in c: bairro = "Calhau"
        elif "TIRIRICAL" in c: bairro = "Tirirical"
        elif "COHATRAC IV" in c: bairro = "Cohatrac IV"
        elif "COHATRAC" in c: bairro = "Cohatrac"
        elif "CAJAZEIRAS" in c: bairro = "Cajazeiras"
        elif "JD" in c: bairro = "Jardim Eldorado"
        elif "BACANGA" in c: bairro = "Bacanga"
        elif "ANJO DA GUARDA" in c: bairro = "Anjo da Guarda"
        elif "RENASCENCA" in c: bairro = "Renascença"
        elif "FABRICA DE PAES" in c: bairro = "Tirirical"
        elif "PANIFICACAO" in c: bairro = "Tirirical"
        elif "MAIOBAO UBATUBA" in c: cidade = "Paço do Lumiar"; bairro = "Maiobão / Ubatuba"
        elif "MAIOBAO" in c: cidade = "Paço do Lumiar"; bairro = "Maiobão"
        elif "MAIOBA" in c: cidade = "Paço do Lumiar"; bairro = "Estrada da Maioba"
        elif "JOAO PAULO" in c: bairro = "João Paulo"
        elif "SAO RAIMUNDO" in c: bairro = "São Raimundo"
        elif "SAO LUIS SHOPPING" in c: bairro = "Jaracaty"
        elif "GENERAL ARTHUR CARVALHO" in c: bairro = "Turu"
        elif "ARACAGY" in c: cidade = "São José de Ribamar"; bairro = "Araçagy"
        elif "FORQUILHA" in c: bairro = "Forquilhas"
        elif "DARK STORE" in c: bairro = "Tirirical"
        elif "SAO CRISTOVAO" in c: bairro = "São Cristóvão"
        elif "DIVINEIA" in c: bairro = "Divinéia"
        elif "PARQUE ATHENAS" in c: bairro = "Parque Athenas"
        elif "SANTA CLARA" in c: bairro = "Santa Clara"
        
    return bairro, cidade, uf

def obter_base_enderecos(df_frios: pd.DataFrame, caminho_enderecos: str) -> dict:
    cadastro_existente = {}
    
    if os.path.exists(caminho_enderecos):
        try:
            print(f"[INFO] Carregando base cadastral de endereços: '{os.path.basename(caminho_enderecos)}'...")
            df_cad = pd.read_excel(caminho_enderecos)
            df_cad.columns = df_cad.columns.astype(str).str.strip()
            
            for _, r in df_cad.iterrows():
                f_cod = formatar_codigo_filial(r.get("Filial", ""))
                if f_cod:
                    cadastro_existente[f_cod] = {
                        "Filial": f_cod,
                        "Nome Loja": str(r.get("Nome Loja", "")).strip(),
                        "Região": str(r.get("Região", "")).strip(),
                        "Rota Padrão": str(r.get("Rota Padrão", "")).strip(),
                        "Restrição Veículo": str(r.get("Restrição Veículo", "Sem restrição")).strip(),
                        "Bairro": str(r.get("Bairro", "")).strip(),
                        "Cidade": str(r.get("Cidade", "")).strip(),
                        "UF": str(r.get("UF", "")).strip(),
                        "Endereço / Referência": str(r.get("Endereço / Referência", "")).strip(),
                        "Latitude": r.get("Latitude", ""),
                        "Longitude": r.get("Longitude", ""),
                        "Observação": str(r.get("Observação", "")).strip()
                    }
        except Exception as e:
            print(f"[AVISO] Falha ao ler base de endereços existente ({e}). Criando nova base...")

    novas_lojas = False
    lojas_operacao = df_frios[["FILIAL_PADRAO", "CLIENTE_PADRAO", "REGIAO_PADRAO", "ROTA_PADRAO"]].drop_duplicates()
    
    for _, r in lojas_operacao.iterrows():
        f_cod = r["FILIAL_PADRAO"]
        if not f_cod:
            continue
        if f_cod not in cadastro_existente:
            novas_lojas = True
            c_nome = r["CLIENTE_PADRAO"]
            regiao = r["REGIAO_PADRAO"]
            bairro, cidade, uf = inferir_dados_localizacao(c_nome, regiao)
            nome_curto = c_nome.split(" - ", 1)[-1] if " - " in c_nome else c_nome
            restr_dict = obter_restricoes_loja(f_cod)
            
            cadastro_existente[f_cod] = {
                "Filial": f_cod,
                "Nome Loja": nome_curto,
                "Região": regiao,
                "Rota Padrão": r["ROTA_PADRAO"],
                "Restrição Veículo": restr_dict["descricao"],
                "Bairro": bairro,
                "Cidade": cidade,
                "UF": uf,
                "Endereço / Referência": f"{bairro}, {cidade} - {uf}" if bairro else f"{cidade} - {uf}",
                "Latitude": "",
                "Longitude": "",
                "Observação": f"[{restr_dict['descricao']}] Cadastrado automaticamente" if restr_dict["descricao"] != "Sem restrição" else "Cadastrado automaticamente pelo sistema"
            }

    if novas_lojas or not os.path.exists(caminho_enderecos):
        try:
            df_salvar = pd.DataFrame(list(cadastro_existente.values()))
            cols = ["Filial", "Nome Loja", "Região", "Rota Padrão", "Restrição Veículo", "Bairro", "Cidade", "UF", "Endereço / Referência", "Latitude", "Longitude", "Observação"]
            cols_finais = [c for c in cols if c in df_salvar.columns]
            df_salvar = df_salvar[cols_finais]
            df_salvar.to_excel(caminho_enderecos, index=False)
            print(f"[OK] Base de endereços salva/atualizada: '{caminho_enderecos}' ({len(df_salvar)} lojas)")
        except Exception as e:
            print(f"[AVISO] Não foi possível atualizar '{caminho_enderecos}' ({e}). Usando dados em memória.")

    return cadastro_existente

def localizar_aba_e_cabecalho(caminho_arquivo: str):
    print(f"[INFO] Inspecionando estrutura das abas em: '{os.path.basename(caminho_arquivo)}'...")
    wb = openpyxl.load_workbook(caminho_arquivo, read_only=True, data_only=True)
    candidatos = []
    
    palavras_setor = ["SECTORES", "SETORES", "SETOR"]
    palavras_rota = ["ROTAS", "ROTA"]
    palavras_peso = ["PESOBRUTO", "PESO BRUTO", "PESO"]
    palavras_regiao = ["REGIONAL", "REGIAO", "REGIÃO"]
    palavras_filial = ["FILIAIS", "FILIAL", "LOJAS", "LOJA"]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row_idx, row in enumerate(ws.iter_rows(max_row=30, values_only=True)):
            linha_texto = " ".join([str(cell).upper().strip() for cell in row if cell is not None])
            
            tem_setor = any(p in linha_texto for p in palavras_setor)
            tem_rota = any(p in linha_texto for p in palavras_rota)
            tem_peso = any(p in linha_texto for p in palavras_peso)
            tem_regiao = any(p in linha_texto for p in palavras_regiao)
            tem_filial = any(p in linha_texto for p in palavras_filial)
            
            pontuacao = sum([tem_setor, tem_rota, tem_peso, tem_regiao, tem_filial])
            if pontuacao >= 4:
                max_rows = ws.max_row or 0
                candidatos.append({
                    "aba": sheet_name,
                    "linha_cabecalho": row_idx,
                    "max_linhas": max_rows,
                    "pontuacao": pontuacao
                })
                break
                
    wb.close()
    
    if not candidatos:
        raise ValueError(
            "Nenhuma aba com as colunas estruturantes (SECTORES, ROTAS, PESOBRUTO, REGIONAL, FILIAIS) foi encontrada."
        )
        
    candidatos.sort(key=lambda c: (c["pontuacao"], c["max_linhas"]), reverse=True)
    melhor_opcao = candidatos[0]
    
    print(f"[OK] Base de dados identificada: Aba '{melhor_opcao['aba']}' ({melhor_opcao['max_linhas']:,} linhas)")
    return melhor_opcao["aba"], melhor_opcao["linha_cabecalho"]

def carregar_dados_frios(caminho_arquivo: str, nome_aba: str, linha_cabecalho: int) -> pd.DataFrame:
    print("[INFO] Carregando dados da aba selecionada...")
    df = pd.read_excel(caminho_arquivo, sheet_name=nome_aba, header=linha_cabecalho)
    df.columns = df.columns.astype(str).str.strip()
    
    col_setor = encontrar_coluna(df.columns, ["Address.SECTORES", "SECTORES", "Address.SETORES", "SETORES", "SETOR"])
    col_rota = encontrar_coluna(df.columns, ["CIclo.ROTAS", "Ciclo.ROTAS", "ROTAS", "CIclo.ROTA", "ROTA"])
    col_regiao = encontrar_coluna(df.columns, ["CIclo.REGIONAL", "Ciclo.REGIONAL", "REGIONAL", "REGIAO", "REGIÃO"])
    col_peso = encontrar_coluna(df.columns, ["PESOBRUTO", "PESO BRUTO", "PESO"])
    col_filial = encontrar_coluna(df.columns, ["FILIAIS", "FILIAL", "LOJAS", "LOJA"])
    col_cliente = encontrar_coluna(df.columns, ["CLIENTE", "CLIENTE2", "NOME", "NOME LOJA"])
    
    colunas_obrigatorias = [col_setor, col_rota, col_regiao, col_peso, col_filial]
    if any(c is None for c in colunas_obrigatorias):
        raise KeyError("Colunas obrigatórias não identificadas na planilha.")
        
    setores_limpos = df[col_setor].astype(str).str.strip().str.upper()
    df_frios = df[setores_limpos.isin([s.upper() for s in SETORES_ALVO])].copy()
    
    if df_frios.empty:
        raise ValueError("Nenhum registro encontrado para os setores CONGELADOS e RESFRIADOS.")
        
    rotas_limpas = df_frios[col_rota].astype(str).str.strip().str.upper()
    mask_rota_valida = (
        df_frios[col_rota].notna() &
        (~rotas_limpas.isin(ROTAS_INVALIDAS)) &
        (rotas_limpas != "")
    )
    df_frios = df_frios[mask_rota_valida].copy()

    df_frios[col_peso] = pd.to_numeric(df_frios[col_peso], errors="coerce").fillna(0)
    df_frios["FILIAL_PADRAO"] = df_frios[col_filial].apply(formatar_codigo_filial)
    df_frios["ROTA_PADRAO"] = df_frios[col_rota].apply(normalizar_nome_rota)
    df_frios["REGIAO_PADRAO"] = df_frios[col_regiao].astype(str).str.strip().str.upper()
    df_frios["CLIENTE_PADRAO"] = df_frios[col_cliente].astype(str).str.strip() if col_cliente else df_frios["FILIAL_PADRAO"]

    col_dias = encontrar_coluna(df.columns, ["Dias Carregamento", "DIAS CARREGAMENTO", "Dias Pedidos", "CARREGAMENTO"])
    if col_dias:
        df_frios["DIAS_CARREGAMENTO"] = df.loc[df_frios.index, col_dias].astype(str).str.strip()
    else:
        df_frios["DIAS_CARREGAMENTO"] = "01-DE HOJE"

    if IGNORAR_LOJAS_4_DIGITOS:
        print("[INFO] Aplicando filtro para ignorar lojas de condomínio de 4 dígitos (preservando Mais Fraldas 14xx e Loja 2061)...")
        mask_4_digitos = df_frios["FILIAL_PADRAO"].apply(
            lambda x: len(x) == 4 and x.isdigit() and not x.startswith("14") and x != "2061"
        )
        total_4_digitos = mask_4_digitos.sum()
        lojas_4_dig_unicas = df_frios.loc[mask_4_digitos, "FILIAL_PADRAO"].nunique()
        df_frios = df_frios[~mask_4_digitos].copy()
        print(f"       {total_4_digitos:,} linhas removidas ({lojas_4_dig_unicas} lojas de condomínio excluídas - 2061 e Mais Fraldas preservadas).")

    print(f"[OK] Registros de Frios elegíveis: {len(df_frios):,}")
    print(f"[OK] Peso bruto total: {df_frios[col_peso].sum():,.2f} kg ({df_frios['FILIAL_PADRAO'].nunique()} lojas)")
    return df_frios

def otimizar_veiculos_rota_pura(rota_nome: str, regiao: str, pool_lojas: list, base_enderecos: dict):
    pool = list(pool_lojas)
    veiculos_dedicados = []

    for it in list(pool):
        if it["filial_orig"] == "96" and it["peso"] > 14000:
            pool.remove(it)
            veiculos_dedicados.append({
                "rota_padrao": rota_nome,
                "regiao": regiao,
                "veiculo": "Truck",
                "peso": 12000.0,
                "filiais": ["96 (Carga 1)"],
                "stores": [{**it, "filial": "96 (Carga 1)", "peso": 12000.0}],
                "tipo_alocacao": "Dedicado Loja 96 (12t)"
            })
            pool.append({
                **it,
                "filial": "96 (Sobra)",
                "peso": it["peso"] - 12000.0
            })

    for s in list(pool):
        if s["peso"] >= 24000:
            v_tipo, _, _ = selecionar_menor_veiculo(s["peso"], [s["filial_orig"]], regiao, base_enderecos)
            if v_tipo is not None:
                veiculos_dedicados.append({
                    "rota_padrao": rota_nome,
                    "regiao": regiao,
                    "veiculo": v_tipo,
                    "peso": s["peso"],
                    "filiais": [s["filial"]],
                    "stores": [s],
                    "tipo_alocacao": "Carga Exclusiva / Dedicada"
                })
                pool.remove(s)

    melhor_veiculos = []
    melhor_sobras = list(pool)
    melhor_score = (-1.0, -1.0)

    def buscar(rem_pool, veics_acum):
        nonlocal melhor_veiculos, melhor_sobras, melhor_score
        p_obrig_acum = sum(sum(s["peso"] for s in v["stores"] if s.get("tem_hoje", True)) for v in veics_acum)
        p_tot_acum = sum(v["peso"] for v in veics_acum)
        score = (p_obrig_acum, p_tot_acum)

        if score > melhor_score:
            melhor_score = score
            melhor_veiculos = list(veics_acum)
            melhor_sobras = list(rem_pool)
        elif score == melhor_score and len(veics_acum) < len(melhor_veiculos):
            melhor_veiculos = list(veics_acum)
            melhor_sobras = list(rem_pool)

        n = len(rem_pool)
        if n == 0:
            return

        indices = list(range(n))
        candidatos = []
        for k in range(min(LIMITE_MAX_LOJAS_POR_ROTA, n), 0, -1):
            for combo in itertools.combinations(indices, k):
                grp = [rem_pool[i] for i in combo]

                if regiao == "CAPITAL" and not any(s.get("tem_hoje", False) for s in grp):
                    continue

                p = sum(s["peso"] for s in grp)
                f_list = [s["filial_orig"] for s in grp]
                v, eh_exc, falta_piso = selecionar_menor_veiculo(p, f_list, regiao, base_enderecos)
                if v is not None:
                    p_obrig = sum(s["peso"] for s in grp if s.get("tem_hoje", True))
                    candidatos.append((grp, v, p, p_obrig, eh_exc, falta_piso, combo))

        candidatos.sort(key=lambda x: (x[3], not x[4], x[2]), reverse=True)

        for grp, v, p, p_obrig, eh_exc, falta_piso, combo in candidatos[:12]:
            novo_rem = [rem_pool[i] for i in indices if i not in combo]
            has_compl = any(not s.get("tem_hoje", True) for s in grp)
            if eh_exc:
                tipo_aloc = f"Exceção Capital (Faltam {falta_piso:,.0f} kg para piso ideal de {v} - gerado para autorização)"
            elif v == "3/4" and p < 5500:
                tipo_aloc = "Rota Pura (Dentro da Meta Operacional 3/4: 5.000 a 5.500 kg)"
            elif has_compl:
                tipo_aloc = "Rota Pura (com complemento de peso)"
            else:
                tipo_aloc = "Rota Pura (100% Pedido do Dia)"

            novo_v = {
                "rota_padrao": rota_nome,
                "regiao": regiao,
                "veiculo": v,
                "peso": p,
                "filiais": [s["filial"] for s in grp],
                "stores": grp,
                "tipo_alocacao": tipo_aloc,
                "eh_excecao": eh_exc,
                "falta_piso": falta_piso
            }
            buscar(novo_rem, veics_acum + [novo_v])

    buscar(pool, [])
    return veiculos_dedicados + melhor_veiculos, melhor_sobras

def processar_alocacao_rotas(df_frios: pd.DataFrame, base_enderecos: dict):
    print(f"\n[INFO] Executando roteirização estrita por ROTA PURA e DIAS CARREGAMENTO (Piso mínimo: {PESO_MINIMO_ROTA:,.0f} kg, Limite: 1 a {LIMITE_MAX_LOJAS_POR_ROTA} lojas)...")
    
    lojas_processadas = []
    col_d = "DIAS_CARREGAMENTO" if "DIAS_CARREGAMENTO" in df_frios.columns else "Dias Carregamento"
    
    for (reg, rota, f), grp in df_frios.groupby(["REGIAO_PADRAO", "ROTA_PADRAO", "FILIAL_PADRAO"]):
        p_total = float(grp["PESOBRUTO"].sum())
        p_hoje = float(grp[grp[col_d].astype(str).str.contains("01-DE HOJE|HOJE", case=False, na=False)]["PESOBRUTO"].sum())
        p_prazo = float(grp[grp[col_d].astype(str).str.contains("02-NO PRAZO|PRAZO", case=False, na=False)]["PESOBRUTO"].sum())
        p_urgente = float(grp[grp[col_d].astype(str).str.contains("03-URGENTE|URGENTE", case=False, na=False)]["PESOBRUTO"].sum())
        
        tem_hoje = (p_hoje > 0)
        if tem_hoje:
            status_dias = "01-DE HOJE"
            tipo_demanda = "Pedido do Dia (Obrigatório)"
            obs_demanda = "Pedido do Dia (Peso agrupado 01-Hoje + Complemento)" if (p_prazo > 0 or p_urgente > 0) else "Pedido do Dia (100% Hoje)"
        elif p_urgente > 0:
            status_dias = "03-URGENTE"
            tipo_demanda = "Complemento de Peso (Urgente)"
            obs_demanda = "Complemento de Peso (03-Urgente)"
        else:
            status_dias = "02-NO PRAZO"
            tipo_demanda = "Complemento de Peso (No Prazo)"
            obs_demanda = "Complemento de Peso (02-No Prazo)"
            
        lojas_processadas.append({
            "regiao": reg,
            "rota": rota,
            "filial": f,
            "filial_orig": f,
            "peso": p_total,
            "peso_hoje": p_hoje,
            "peso_prazo": p_prazo,
            "peso_urgente": p_urgente,
            "tem_hoje": tem_hoje,
            "status_dias": status_dias,
            "tipo_demanda": tipo_demanda,
            "obs_demanda": obs_demanda
        })
        
    df_lojas_proc = pd.DataFrame(lojas_processadas)
    rotas_aprovadas = []
    detalhes_lojas = []
    todas_sobras = []
    
    pares_rotas = (
        df_lojas_proc[["regiao", "rota"]]
        .drop_duplicates()
        .copy()
    )
    pares_rotas["ordem_reg"] = pares_rotas["regiao"].apply(lambda r: 0 if "CAPITAL" in str(r).upper() else 1)
    pares_rotas["num_rota"] = pares_rotas["rota"].apply(lambda r: int(''.join(filter(str.isdigit, str(r)))) if any(c.isdigit() for c in str(r)) else 999)
    pares_rotas = pares_rotas.sort_values(by=["ordem_reg", "num_rota", "rota"])
    
    for _, r_par in pares_rotas.iterrows():
        reg = r_par["regiao"]
        r_nome = r_par["rota"]
        sub = df_lojas_proc[(df_lojas_proc["regiao"] == reg) & (df_lojas_proc["rota"] == r_nome)]
        lojas_rota = sub.to_dict("records")
        
        veics, sobras = otimizar_veiculos_rota_pura(r_nome, reg, lojas_rota, base_enderecos)
        
        for idx_v, v in enumerate(veics, 1):
            r_exec = f"{r_nome} - Veículo {idx_v}"
            v["rota_execucao"] = r_exec
            rotas_aprovadas.append(v)
            for s in v["stores"]:
                detalhes_lojas.append({
                    "regiao": v["regiao"],
                    "rota_padrao": s["rota"],
                    "rota_execucao": r_exec,
                    "veiculo": v["veiculo"],
                    "filial": s["filial"],
                    "peso": s["peso"],
                    "peso_hoje": s.get("peso_hoje", s["peso"]),
                    "peso_prazo": s.get("peso_prazo", 0.0),
                    "peso_urgente": s.get("peso_urgente", 0.0),
                    "tem_hoje": s.get("tem_hoje", True),
                    "status_dias": s.get("status_dias", "01-DE HOJE"),
                    "tipo_demanda": s.get("tipo_demanda", "Pedido do Dia (Obrigatório)"),
                    "obs_demanda": s.get("obs_demanda", "")
                })
                
        for s in sobras:
            todas_sobras.append(s)

    cargas_aguardando = []
    peso_por_rota_sobra = Counter()
    for s in todas_sobras:
        peso_por_rota_sobra[s["rota"]] += s["peso"]

    for s in todas_sobras:
        f_cod = s["filial_orig"]
        info = base_enderecos.get(f_cod, {})
        restr_texto = info.get("Restrição Veículo", "Sem restrição")
        valvo = obter_menor_veiculo_alvo([f_cod], base_enderecos, s["regiao"])
        p_acum_rota = peso_por_rota_sobra[s["rota"]]
        falta_piso = max(0.0, valvo["peso_minimo"] - p_acum_rota)
        
        if s.get("tem_hoje", False):
            motivo = (
                f"Pedido do Dia (01-DE HOJE) aguardando geração de pedidos adicionais na Rota {s['rota']}. "
                f"Possui {p_acum_rota:,.2f} kg acumulados; faltam {falta_piso:,.2f} kg para atingir o piso mínimo do {valvo['tipo']} ({valvo['peso_minimo']:,.0f} kg) sem quebra de frete."
            )
        else:
            motivo = (
                f"Carga com status {s.get('status_dias', '02-NO PRAZO')}. Atua como complemento de peso para pedidos do dia na Rota {s['rota']} (aguardando liberação de rota)."
            )
        
        cargas_aguardando.append({
            "Região": s["regiao"],
            "Rota Padrão": s["rota"],
            "Filial": s["filial"],
            "Nome Loja": info.get("Nome Loja", ""),
            "Dias Carregamento": s.get("status_dias", "02-NO PRAZO"),
            "Tipo Demanda": s.get("tipo_demanda", "Complemento de Peso (No Prazo)"),
            "Restrição Veículo": restr_texto,
            "Bairro": info.get("Bairro", ""),
            "Cidade": info.get("Cidade", ""),
            "UF": info.get("UF", ""),
            "Endereço / Referência": info.get("Endereço / Referência", ""),
            "Latitude": info.get("Latitude", ""),
            "Longitude": info.get("Longitude", ""),
            "Peso Loja (Kg)": round(s["peso"], 2),
            "Peso Acumulado Rota (Kg)": round(p_acum_rota, 2),
            "Veículo Alvo": valvo["tipo"],
            "Piso Mínimo Veículo (Kg)": valvo["peso_minimo"],
            "Falta para Bater Piso (Kg)": round(falta_piso, 2),
            "Status / Motivo": motivo
        })

    rotas_finais_formatadas = []
    for r in rotas_aprovadas:
        paradas_ordenadas = ordenar_paradas(r["filiais"], r["regiao"], base_enderecos)
        p1 = paradas_ordenadas[0] if len(paradas_ordenadas) > 0 else ""
        p2 = paradas_ordenadas[1] if len(paradas_ordenadas) > 1 else ""
        p3 = paradas_ordenadas[2] if len(paradas_ordenadas) > 2 else ""
        p4 = paradas_ordenadas[3] if len(paradas_ordenadas) > 3 else ""
        
        obs_parts = []
        for f in paradas_ordenadas:
            fc = str(f).split()[0].strip()
            restr = base_enderecos.get(fc, {}).get("Restrição Veículo")
            if restr and restr != "Sem restrição":
                obs_parts.append(f"Loja {fc}: {restr}")
        if "Dedicado Loja 96" in r.get("tipo_alocacao", ""):
            obs_parts.append("Carga dedicada 12t fechada para Loja 96")
        else:
            obs_parts.append(r.get("tipo_alocacao", f"Rota Pura {r['rota_padrao']}"))
            
        locais = []
        for f in paradas_ordenadas:
            fc = str(f).split()[0].strip()
            info_f = base_enderecos.get(fc, {})
            b = info_f.get("Bairro", "")
            cid = info_f.get("Cidade", "")
            loc = f"{b} ({cid})" if b else cid
            if loc and loc not in locais:
                locais.append(loc)
                
        rotas_finais_formatadas.append({
            "Região": r["regiao"],
            "Rota Padrão": r["rota_padrao"],
            "Rota de Execução": r["rota_execucao"],
            "Veículo Ideal": r["veiculo"],
            "Peso Total (Kg)": round(r["peso"], 2),
            "1ª Entrega": p1,
            "2ª Entrega": p2,
            "3ª Entrega": p3,
            "4ª Entrega": p4,
            "Filiais Atendidas": " | ".join(paradas_ordenadas),
            "Bairros / Cidades Atendidos": " | ".join(locais),
            "Qtd Filiais": len(paradas_ordenadas),
            "Separação": "AMBOS",
            "Observação": " | ".join(obs_parts)
        })

    detalhes_finais = []
    for d in detalhes_lojas:
        fc = str(d["filial"]).split()[0].strip()
        inf = base_enderecos.get(fc, {})
        detalhes_finais.append({
            "Região": d["regiao"],
            "Rota Padrão": d["rota_padrao"],
            "Rota de Execução": d["rota_execucao"],
            "Veículo Ideal": d["veiculo"],
            "Filial": d["filial"],
            "Nome Loja": inf.get("Nome Loja", ""),
            "Dias Carregamento": d.get("status_dias", "01-DE HOJE"),
            "Tipo Demanda": d.get("tipo_demanda", "Pedido do Dia (Obrigatório)"),
            "Restrição Veículo": inf.get("Restrição Veículo", "Sem restrição"),
            "Bairro": inf.get("Bairro", ""),
            "Cidade": inf.get("Cidade", ""),
            "UF": inf.get("UF", ""),
            "Endereço / Referência": inf.get("Endereço / Referência", ""),
            "Latitude": inf.get("Latitude", ""),
            "Longitude": inf.get("Longitude", ""),
            "Peso Hoje (Kg)": round(d.get("peso_hoje", d["peso"]), 2),
            "Peso Complemento (Kg)": round(d.get("peso_prazo", 0.0) + d.get("peso_urgente", 0.0), 2),
            "Peso Loja (Kg)": round(d["peso"], 2),
            "Observação Demanda": d.get("obs_demanda", "")
        })

    df_resultado = pd.DataFrame(rotas_finais_formatadas)
    df_detalhes = pd.DataFrame(detalhes_finais)
    df_aguardando = pd.DataFrame(cargas_aguardando)
    df_sobras = pd.DataFrame()

    return df_resultado, df_detalhes, df_aguardando, df_sobras

def exportar_para_excel_formatado(
    df_resumo: pd.DataFrame,
    df_detalhes: pd.DataFrame,
    df_aguardando: pd.DataFrame,
    df_sobras: pd.DataFrame,
    caminho_saida: str,
    base_enderecos: dict = None
):
    print(f"\n[INFO] Formatando e gravando planilha executiva em: '{os.path.basename(caminho_saida)}'...")
    wb = openpyxl.Workbook()
    
    cor_azul_escuro = "1B365D"
    cor_azul_sub = "2B4C7E"
    cor_petroleo = "0D5C75"
    cor_ambar = "C65911"
    cor_vinho = "8B1E1E"
    cor_zebra = "F9FAFB"
    cor_light = "F2F4F7"
    
    font_top_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_header_table = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, bold=False, color="1F2937")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color="000000")
    
    fill_dark_blue = PatternFill(start_color=cor_azul_escuro, end_color=cor_azul_escuro, fill_type="solid")
    fill_navy_sub = PatternFill(start_color=cor_azul_sub, end_color=cor_azul_sub, fill_type="solid")
    fill_petroleo = PatternFill(start_color=cor_petroleo, end_color=cor_petroleo, fill_type="solid")
    fill_ambar = PatternFill(start_color=cor_ambar, end_color=cor_ambar, fill_type="solid")
    fill_vinho = PatternFill(start_color=cor_vinho, end_color=cor_vinho, fill_type="solid")
    fill_light_gray = PatternFill(start_color=cor_light, end_color=cor_light, fill_type="solid")
    fill_zebra = PatternFill(start_color=cor_zebra, end_color=cor_zebra, fill_type="solid")
    fill_zebra_ambar = PatternFill(start_color="FFF8F0", end_color="FFF8F0", fill_type="solid")
    fill_total = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='D0D5DD'),
        right=Side(style='thin', color='D0D5DD'),
        top=Side(style='thin', color='D0D5DD'),
        bottom=Side(style='thin', color='D0D5DD')
    )
    border_total = Border(
        top=Side(style='thin', color='000000'),
        bottom=Side(style='double', color='000000'),
        left=Side(style='thin', color='D0D5DD'),
        right=Side(style='thin', color='D0D5DD')
    )
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    def criar_aba_operacional(ws, titulo, df_regiao_rotas):
        ws.title = titulo
        ws.views.sheetView[0].showGridLines = True
        
        tipos_frota = ["3/4.", "TOCO", "TRUCK", "BITRUCK", "CARRETA"]
        cont_v = Counter()
        for _, r in df_regiao_rotas.iterrows():
            v = str(r["Veículo Ideal"]).upper().replace(".", "").strip()
            if "3/4" in v: cont_v["3/4."] += 1
            elif "TOCO" in v: cont_v["TOCO"] += 1
            elif "TRUCK" in v: cont_v["TRUCK"] += 1
            elif "BITRUCK" in v: cont_v["BITRUCK"] += 1
            elif "CARRETA" in v: cont_v["CARRETA"] += 1
            
        for c_idx, t in enumerate(tipos_frota, 2):
            c1 = ws.cell(row=1, column=c_idx, value=t)
            c1.font = font_top_header; c1.fill = fill_dark_blue; c1.alignment = align_center; c1.border = border_thin
            c2 = ws.cell(row=2, column=c_idx, value=cont_v[t])
            c2.font = font_bold; c2.alignment = align_center; c2.fill = fill_light_gray; c2.border = border_thin
            
        ws.cell(row=4, column=2, value="TOTAL VEÍCULOS").font = font_top_header
        ws.cell(row=4, column=2).fill = fill_navy_sub; ws.cell(row=4, column=2).alignment = align_center; ws.cell(row=4, column=2).border = border_thin
        
        rot_ton_label = "TON. TOTAIS CAP" if "CAP" in titulo else "TON. GERADAS"
        ws.cell(row=4, column=4, value=rot_ton_label).font = font_top_header
        ws.cell(row=4, column=4).fill = fill_navy_sub; ws.cell(row=4, column=4).alignment = align_center; ws.cell(row=4, column=4).border = border_thin
        
        c_tot_v = ws.cell(row=5, column=2, value=len(df_regiao_rotas))
        c_tot_v.font = font_bold; c_tot_v.alignment = align_center; c_tot_v.fill = fill_light_gray; c_tot_v.border = border_thin
        
        peso_tot = df_regiao_rotas["Peso Total (Kg)"].sum() if not df_regiao_rotas.empty else 0.0
        c_tot_p = ws.cell(row=5, column=4, value=round(peso_tot, 2))
        c_tot_p.font = font_bold; c_tot_p.alignment = align_right; c_tot_p.number_format = '#,##0.00'; c_tot_p.fill = fill_light_gray; c_tot_p.border = border_thin
        
        headers = ["1ª ENTREGA", "2ª ENTREGA", "3ª ENTREGA", "4ª ENTREGA", "SEPARAÇÃO", "VEICULO", "PESO TOTAL (KG)", "ROTA PADRÃO", "OBSERVAÇÃO"]
        for c_idx, h in enumerate(headers, 2):
            c = ws.cell(row=7, column=c_idx, value=h)
            c.font = font_header_table; c.fill = fill_dark_blue; c.alignment = align_center; c.border = border_thin
            ws.row_dimensions[7].height = 26
            
        if not df_regiao_rotas.empty:
            for r_idx, (_, r) in enumerate(df_regiao_rotas.iterrows(), 8):
                ws.row_dimensions[r_idx].height = 20
                is_zebra = (r_idx % 2 == 0)
                
                vals = [
                    r.get("1ª Entrega", ""),
                    r.get("2ª Entrega", ""),
                    r.get("3ª Entrega", ""),
                    r.get("4ª Entrega", ""),
                    r.get("Separação", "AMBOS"),
                    r.get("Veículo Ideal", ""),
                    r.get("Peso Total (Kg)", 0.0),
                    r.get("Rota Padrão", ""),
                    r.get("Observação", "")
                ]
                
                for c_idx, val in enumerate(vals, 2):
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.font = font_data; cell.border = border_thin
                    if is_zebra: cell.fill = fill_zebra
                    
                    if c_idx in [2, 3, 4, 5]: cell.alignment = align_center
                    elif c_idx == 6: cell.alignment = align_center
                    elif c_idx == 7: cell.alignment = align_center; cell.font = font_bold
                    elif c_idx == 8: cell.alignment = align_right; cell.number_format = '#,##0.00'
                    elif c_idx == 9: cell.alignment = align_center; cell.font = font_bold
                    else: cell.alignment = align_left
        else:
            ws.cell(row=8, column=2, value="Nenhuma rota nesta região.").font = font_data
            
        for col_idx in range(2, 11):
            col_letter = get_column_letter(col_idx)
            max_l = max(len(str(ws.cell(r, col_idx).value or '')) for r in range(1, len(df_regiao_rotas) + 9))
            ws.column_dimensions[col_letter].width = max(max_l + 3, 14)
        ws.column_dimensions["J"].width = 50

    df_cap = df_resumo[df_resumo["Região"] == "CAPITAL"].copy() if not df_resumo.empty else pd.DataFrame()
    ws_cap = wb.active
    criar_aba_operacional(ws_cap, "ROTAS_CAP", df_cap)
    
    df_int = df_resumo[df_resumo["Região"] == "INTERIOR"].copy() if not df_resumo.empty else pd.DataFrame()
    ws_int = wb.create_sheet()
    criar_aba_operacional(ws_int, "ROTA_INT", df_int)

    ws3 = wb.create_sheet(title="Roteirização Consolidada")
    ws3.views.sheetView[0].showGridLines = True
    
    cols_exec = ["Região", "Rota Padrão", "Rota de Execução", "Veículo Ideal", "Peso Total (Kg)", "1ª Entrega", "2ª Entrega", "3ª Entrega", "4ª Entrega", "Bairros / Cidades Atendidos", "Qtd Filiais"]
    cols_existentes = [c for c in cols_exec if c in df_resumo.columns] if not df_resumo.empty else cols_exec
    
    for c_idx, col_nome in enumerate(cols_existentes, 1):
        cell = ws3.cell(row=1, column=c_idx, value=col_nome)
        cell.font = font_header_table; cell.fill = fill_dark_blue; cell.alignment = align_center; cell.border = border_thin
        ws3.row_dimensions[1].height = 28
        
    if not df_resumo.empty:
        for r_idx, (_, r_data) in enumerate(df_resumo[cols_existentes].iterrows(), 2):
            ws3.row_dimensions[r_idx].height = 20
            is_zebra = (r_idx % 2 == 0)
            for c_idx, col_nome in enumerate(cols_existentes, 1):
                cell = ws3.cell(row=r_idx, column=c_idx, value=r_data[col_nome])
                cell.font = font_data; cell.border = border_thin
                if is_zebra: cell.fill = fill_zebra
                
                if col_nome in ["Região", "Rota Padrão", "Veículo Ideal", "1ª Entrega", "2ª Entrega", "3ª Entrega", "4ª Entrega"]:
                    cell.alignment = align_center
                elif col_nome == "Peso Total (Kg)":
                    cell.alignment = align_right; cell.number_format = '#,##0.00'
                elif col_nome == "Qtd Filiais":
                    cell.alignment = align_center; cell.number_format = '#,##0'
                else:
                    cell.alignment = align_left
                    
        total_row3 = len(df_resumo) + 2
        ws3.row_dimensions[total_row3].height = 24
        for c in range(1, len(cols_existentes) + 1):
            cell = ws3.cell(row=total_row3, column=c)
            cell.font = font_total; cell.fill = fill_total; cell.border = border_total
            
        ws3.cell(row=total_row3, column=1, value="TOTAL GERAL").alignment = align_center
        if "Peso Total (Kg)" in cols_existentes:
            c_p_idx = cols_existentes.index("Peso Total (Kg)") + 1
            l_p = get_column_letter(c_p_idx)
            cell_p = ws3.cell(row=total_row3, column=c_p_idx, value=f"=SUM({l_p}2:{l_p}{total_row3-1})")
            cell_p.alignment = align_right; cell_p.number_format = '#,##0.00'
    else:
        ws3.cell(row=2, column=1, value="Nenhuma rota gerada.").font = font_data
        total_row3 = 3
        
    for col in ws3.columns:
        col_letter = get_column_letter(col[0].column)
        max_l = max(len(str(cell.value or '')) for cell in col if not str(cell.value).startswith('='))
        largura = max(max_l + 3, 13)
        if "Bairros / Cidades Atendidos" in cols_existentes and col_letter == get_column_letter(cols_existentes.index("Bairros / Cidades Atendidos") + 1):
            largura = min(largura, 50)
        ws3.column_dimensions[col_letter].width = largura
    ws3.freeze_panes = "A2"

    ws4 = wb.create_sheet(title="Detalhe Lojas e Endereços")
    ws4.views.sheetView[0].showGridLines = True
    colunas4 = list(df_detalhes.columns) if not df_detalhes.empty else [
        "Região", "Rota Padrão", "Rota de Execução", "Veículo Ideal", "Filial", "Nome Loja", "Dias Carregamento", "Tipo Demanda", "Restrição Veículo", "Bairro", "Cidade", "UF", "Endereço / Referência", "Latitude", "Longitude", "Peso Hoje (Kg)", "Peso Complemento (Kg)", "Peso Loja (Kg)", "Observação Demanda"
    ]
    
    for c_idx, col_nome in enumerate(colunas4, 1):
        cell = ws4.cell(row=1, column=c_idx, value=col_nome)
        cell.font = font_header_table; cell.fill = fill_petroleo; cell.alignment = align_center; cell.border = border_thin
        ws4.row_dimensions[1].height = 28
        
    if not df_detalhes.empty:
        for r_idx, (_, r_data) in enumerate(df_detalhes.iterrows(), 2):
            ws4.row_dimensions[r_idx].height = 20
            is_zebra = (r_idx % 2 == 0)
            for c_idx, col_nome in enumerate(colunas4, 1):
                valor = r_data[col_nome]
                cell = ws4.cell(row=r_idx, column=c_idx, value=valor)
                cell.font = font_data; cell.border = border_thin
                if is_zebra: cell.fill = fill_zebra
                
                if col_nome in ["Região", "Rota Padrão", "Veículo Ideal", "Filial", "UF", "Restrição Veículo", "Dias Carregamento", "Tipo Demanda"]:
                    cell.alignment = align_center
                elif col_nome in ["Latitude", "Longitude"]:
                    cell.alignment = align_center
                elif col_nome in ["Peso Hoje (Kg)", "Peso Complemento (Kg)", "Peso Loja (Kg)"]:
                    cell.alignment = align_right; cell.number_format = '#,##0.00'
                else:
                    cell.alignment = align_left
                    
                if col_nome == "Restrição Veículo" and str(valor).strip() not in ("", "Sem restrição"):
                    cell.font = Font(name="Calibri", size=10, bold=True, color="9C0006")
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                    
        total_row4 = len(df_detalhes) + 2
        ws4.row_dimensions[total_row4].height = 24
        for c in range(1, len(colunas4) + 1):
            cell = ws4.cell(row=total_row4, column=c)
            cell.font = font_total; cell.fill = fill_total; cell.border = border_total
            
        ws4.cell(row=total_row4, column=1, value="TOTAL GERAL").alignment = align_center
        for col_p_nome in ["Peso Hoje (Kg)", "Peso Complemento (Kg)", "Peso Loja (Kg)"]:
            if col_p_nome in colunas4:
                col_p_idx4 = colunas4.index(col_p_nome) + 1
                l_p4 = get_column_letter(col_p_idx4)
                cell_p4 = ws4.cell(row=total_row4, column=col_p_idx4, value=f"=SUM({l_p4}2:{l_p4}{total_row4-1})")
                cell_p4.alignment = align_right; cell_p4.number_format = '#,##0.00'
    else:
        ws4.cell(row=2, column=1, value="Nenhuma loja roteirizada.").font = font_data
        total_row4 = 3
        
    for col in ws4.columns:
        col_letter = get_column_letter(col[0].column)
        max_l = max(len(str(cell.value or '')) for cell in col if not str(cell.value).startswith('='))
        largura = max(max_l + 3, 12)
        if "Nome Loja" in colunas4 and col_letter == get_column_letter(colunas4.index("Nome Loja") + 1):
            largura = min(largura, 45)
        elif "Endereço / Referência" in colunas4 and col_letter == get_column_letter(colunas4.index("Endereço / Referência") + 1):
            largura = min(largura, 45)
        elif "Restrição Veículo" in colunas4 and col_letter == get_column_letter(colunas4.index("Restrição Veículo") + 1):
            largura = max(largura, 22)
        elif "Observação Demanda" in colunas4 and col_letter == get_column_letter(colunas4.index("Observação Demanda") + 1):
            largura = min(largura, 45)
        ws4.column_dimensions[col_letter].width = largura
    ws4.freeze_panes = "A2"

    ws5 = wb.create_sheet(title="Aguardando Geração (Bater Peso)")
    ws5.views.sheetView[0].showGridLines = True
    cols5 = [
        "Região", "Rota Padrão", "Filial", "Nome Loja", "Dias Carregamento", "Tipo Demanda", "Restrição Veículo", "Bairro", "Cidade", "UF",
        "Endereço / Referência", "Latitude", "Longitude", "Peso Loja (Kg)", "Peso Acumulado Rota (Kg)",
        "Veículo Alvo", "Piso Mínimo Veículo (Kg)", "Falta para Bater Piso (Kg)", "Status / Motivo"
    ]
    cols_existentes5 = [c for c in cols5 if c in df_aguardando.columns] if not df_aguardando.empty else cols5

    for c_idx, col_nome in enumerate(cols_existentes5, 1):
        cell = ws5.cell(row=1, column=c_idx, value=col_nome)
        cell.font = font_header_table; cell.fill = fill_ambar; cell.alignment = align_center; cell.border = border_thin
        ws5.row_dimensions[1].height = 28

    if not df_aguardando.empty:
        for r_idx, (_, r_data) in enumerate(df_aguardando[cols_existentes5].iterrows(), 2):
            ws5.row_dimensions[r_idx].height = 20
            is_zebra = (r_idx % 2 == 0)
            for c_idx, col_nome in enumerate(cols_existentes5, 1):
                valor = r_data[col_nome]
                cell = ws5.cell(row=r_idx, column=c_idx, value=valor)
                cell.font = font_data; cell.border = border_thin
                if is_zebra: cell.fill = fill_zebra_ambar
                if col_nome in ["Região", "Rota Padrão", "Filial", "UF", "Restrição Veículo", "Veículo Alvo", "Dias Carregamento", "Tipo Demanda"]:
                    cell.alignment = align_center
                elif col_nome in ["Latitude", "Longitude"]:
                    cell.alignment = align_center
                elif col_nome in ["Peso Loja (Kg)", "Peso Acumulado Rota (Kg)", "Piso Mínimo Veículo (Kg)"]:
                    cell.alignment = align_right; cell.number_format = '#,##0.00'
                elif col_nome == "Falta para Bater Piso (Kg)":
                    cell.alignment = align_right; cell.number_format = '#,##0.00'
                    cell.font = font_bold
                else:
                    cell.alignment = align_left
                if col_nome == "Restrição Veículo" and str(valor).strip() not in ("", "Sem restrição"):
                    cell.font = Font(name="Calibri", size=10, bold=True, color="9C0006")
                    cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        tot_r5 = len(df_aguardando) + 2
        ws5.row_dimensions[tot_r5].height = 24
        for c in range(1, len(cols_existentes5) + 1):
            cell = ws5.cell(row=tot_r5, column=c)
            cell.font = font_total; cell.fill = fill_total; cell.border = border_total
        ws5.cell(row=tot_r5, column=1, value="TOTAL GERAL").alignment = align_center
        if "Peso Loja (Kg)" in cols_existentes5:
            col_p_idx5 = cols_existentes5.index("Peso Loja (Kg)") + 1
            l_p5 = get_column_letter(col_p_idx5)
            cell_p5 = ws5.cell(row=tot_r5, column=col_p_idx5, value=f"=SUM({l_p5}2:{l_p5}{tot_r5-1})")
            cell_p5.alignment = align_right; cell_p5.number_format = '#,##0.00'
    else:
        ws5.cell(row=2, column=1, value="Nenhuma carga aguardando geração nesta operação.").font = font_data

    for col in ws5.columns:
        col_letter = get_column_letter(col[0].column)
        max_l = max(len(str(cell.value or '')) for cell in col if not str(cell.value).startswith('='))
        ws5.column_dimensions[col_letter].width = max(max_l + 3, 14)
    if "Status / Motivo" in cols_existentes5:
        col_m_letter = get_column_letter(cols_existentes5.index("Status / Motivo") + 1)
        ws5.column_dimensions[col_m_letter].width = 65
    ws5.freeze_panes = "A2"

    ws6 = wb.create_sheet(title="Sobras (Não Roteirizadas)")
    ws6.views.sheetView[0].showGridLines = True
    colunas6 = list(df_sobras.columns) if not df_sobras.empty else [
        "Região", "Rota Padrão", "Filial", "Nome Loja", "Restrição Veículo", "Bairro", "Cidade", "UF",
        "Endereço / Referência", "Latitude", "Longitude", "Peso Loja (Kg)", "Status / Motivo"
    ]
    for c_idx, col_nome in enumerate(colunas6, 1):
        cell = ws6.cell(row=1, column=c_idx, value=col_nome)
        cell.font = font_header_table; cell.fill = fill_vinho; cell.alignment = align_center; cell.border = border_thin
        ws6.row_dimensions[1].height = 28

    if not df_sobras.empty:
        for r_idx, (_, r_data) in enumerate(df_sobras.iterrows(), 2):
            ws6.row_dimensions[r_idx].height = 20
            is_zebra = (r_idx % 2 == 0)
            for c_idx, col_nome in enumerate(colunas6, 1):
                valor = r_data[col_nome]
                cell = ws6.cell(row=r_idx, column=c_idx, value=valor)
                cell.font = font_data; cell.border = border_thin
                if is_zebra: cell.fill = PatternFill(start_color="FFF5F5", end_color="FFF5F5", fill_type="solid")
                if col_nome in ["Região", "Rota Padrão", "Filial", "UF", "Restrição Veículo"]:
                    cell.alignment = align_center
                elif col_nome in ["Latitude", "Longitude"]:
                    cell.alignment = align_center
                elif col_nome == "Peso Loja (Kg)":
                    cell.alignment = align_right; cell.number_format = '#,##0.00'
                else:
                    cell.alignment = align_left
        tot_r6 = len(df_sobras) + 2
        ws6.row_dimensions[tot_r6].height = 24
        for c in range(1, len(colunas6) + 1):
            cell = ws6.cell(row=tot_r6, column=c)
            cell.font = font_total; cell.fill = fill_total; cell.border = border_total
        ws6.cell(row=tot_r6, column=1, value="TOTAL GERAL").alignment = align_center
    else:
        ws6.cell(row=2, column=1, value="Nenhuma sobra técnica nesta operação (todas as cargas pendentes estão na aba 'Aguardando Geração (Bater Peso)').").font = font_data

    for col in ws6.columns:
        col_letter = get_column_letter(col[0].column)
        ws6.column_dimensions[col_letter].width = 25
    ws6.freeze_panes = "A2"

    try:
        wb.save(caminho_saida)
        print(f"[OK] Planilha executiva com 6 abas salva com sucesso: '{caminho_saida}'")
    except PermissionError:
        nome_alt = caminho_saida.replace(".xlsx", "_Novo.xlsx")
        try:
            wb.save(nome_alt)
            print(f"\n[AVISO] '{os.path.basename(caminho_saida)}' está atualmente aberto no Excel!")
            print(f"[OK] Planilha executiva salva com sucesso em: '{nome_alt}'")
        except Exception:
            raise

def executar_roteirizacao(
    caminho_entrada: str = ARQUIVO_ENTRADA_PADRAO,
    caminho_saida: str = ARQUIVO_SAIDA_PADRAO,
    caminho_enderecos: str = ARQUIVO_ENDERECOS_PADRAO
):
    print("=" * 80)
    print("INICIANDO PROCESSAMENTO DE ROTEIRIZAÇÃO - SETOR DE FRIOS (CALIBRADO)")
    print("=" * 80)
    
    if not os.path.exists(caminho_entrada):
        print(f"\n[ERRO] Arquivo de entrada não encontrado: '{caminho_entrada}'")
        return False
        
    try:

        aba_correta, linha_cabecalho = localizar_aba_e_cabecalho(caminho_entrada)
        
        df_frios = carregar_dados_frios(caminho_entrada, aba_correta, linha_cabecalho)
        
        base_enderecos = obter_base_enderecos(df_frios, caminho_enderecos)
        
        for idx in df_frios.index:
            f = df_frios.at[idx, "FILIAL_PADRAO"]
            if f in base_enderecos:
                info_cad = base_enderecos[f]
                r_cad = info_cad.get("Rota Padrão")
                reg_cad = info_cad.get("Região")
                if r_cad and str(r_cad).strip() not in ("", "nan", "None"):
                    df_frios.at[idx, "ROTA_PADRAO"] = normalizar_nome_rota(r_cad)
                if reg_cad and str(reg_cad).strip() not in ("", "nan", "None"):
                    df_frios.at[idx, "REGIAO_PADRAO"] = str(reg_cad).strip().upper()
        
        df_resumo, df_detalhes, df_aguardando, df_sobras = processar_alocacao_rotas(df_frios, base_enderecos)
        
        exportar_para_excel_formatado(df_resumo, df_detalhes, df_aguardando, df_sobras, caminho_saida, base_enderecos)
        
        print("\n" + "=" * 80)
        print("RESUMO DA ROTEIRIZAÇÃO CONCLUÍDA (FAIXAS ESTRITAS DE PESO)")
        print("=" * 80)
        if not df_resumo.empty:
            cols_print = ["Região", "Rota de Execução", "Veículo Ideal", "Peso Total (Kg)", "1ª Entrega", "2ª Entrega", "3ª Entrega", "4ª Entrega"]
            cols_print = [c for c in cols_print if c in df_resumo.columns]
            print(df_resumo[cols_print].to_string(index=False))
            print("=" * 80)
            print(f"[RESUMO] Total de Veículos/Rotas Aprovados: {len(df_resumo)}")
            print(f"[RESUMO] Peso Total Roteirizado: {df_resumo['Peso Total (Kg)'].sum():,.2f} kg")
            print(f"[RESUMO] Total de Lojas Roteirizadas: {df_resumo['Qtd Filiais'].sum():.0f} paradas")
            print(f"[RESUMO] Distribuição de Lojas por Veículo:")
            dist = df_resumo["Qtd Filiais"].value_counts().sort_index()
            for qtd, count in dist.items():
                print(f"         • Veículos com {int(qtd)} loja(s): {count}")
        else:
            print("[AVISO] Nenhuma rota atingiu o peso mínimo operacional.")

        if not df_aguardando.empty:
            peso_aguardando = df_aguardando["Peso Loja (Kg)"].sum()
            lojas_aguardando = df_aguardando["Filial"].nunique()
            rotas_aguardando = df_aguardando["Rota Padrão"].nunique()
            falta_total = df_aguardando.groupby("Rota Padrão")["Falta para Bater Piso (Kg)"].first().sum()
            print("-" * 80)
            print(f"[AGUARDANDO GERAÇÃO] Total de Cargas Pendentes de Piso: {peso_aguardando:,.2f} kg ({lojas_aguardando} lojas em {rotas_aguardando} rotas)")
            print(f"[AGUARDANDO GERAÇÃO] Peso Faltante para Bater Piso Operacional: {falta_total:,.2f} kg (evita quebra de meta/queima de frete)")
            print(f"[AGUARDANDO GERAÇÃO] Detalhes disponíveis na aba 'Aguardando Geração (Bater Peso)' do Excel.")

        if not df_sobras.empty:
            peso_sobras = df_sobras["Peso Loja (Kg)"].sum()
            lojas_sobras = df_sobras["Filial"].nunique()
            rotas_sobras = df_sobras["Rota Padrão"].nunique()
            print("-" * 80)
            print(f"[SOBRAS TÉCNICAS] Cargas Não Roteirizadas: {peso_sobras:,.2f} kg ({lojas_sobras} lojas)")
            print(f"[SOBRAS TÉCNICAS] Detalhes disponíveis na aba 'Sobras (Não Roteirizadas)' do Excel.")

        print(f"[RESUMO] Arquivo consolidado: {caminho_saida}")
        print(f"[RESUMO] Arquivo cadastral de endereços: {caminho_enderecos}")
        print("=" * 80)
        return True

    except PermissionError as e:
        print("\n" + "!" * 80)
        print("[ERRO DE PERMISSÃO] ARQUIVO EM USO")
        print("!" * 80)
        print(f"Detalhe: {e}")
        print("SOLUÇÃO: Feche o arquivo Excel ('Acomp_EvoluSep.xlsx', 'Prototipo_Rotas_Frios.xlsx' ou 'Cadastro_Lojas_Enderecos.xlsx')")
        print("         e tente executar novamente o script.")
        print("!" * 80)
        return False

    except (zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as e:
        print("\n" + "!" * 80)
        print("[ERRO] ARQUIVO EXCEL CORROMPIDO OU INVÁLIDO")
        print("!" * 80)
        print(f"Detalhe: {e}")
        print("SOLUÇÃO: Abra o arquivo no Excel, salve-o novamente e execute o script.")
        print("!" * 80)
        return False

    except Exception as e:
        print("\n" + "!" * 80)
        print(f"[ERRO INESPERADO] {type(e).__name__} - {e}")
        import traceback
        traceback.print_exc()
        print("!" * 80)
        return False

if __name__ == "__main__":
    executar_roteirizacao()
