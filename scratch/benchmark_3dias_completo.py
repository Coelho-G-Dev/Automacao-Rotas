# -*- coding: utf-8 -*-
"""
BENCHMARK COMPLETO DE 3 DIAS REAIS (24/09, 25/09, 26/09)
Compara as rotas reais montadas manualmente pelos operadores com as rotas
geradas pelo algoritmo automatizado, avaliando aderência, veículos e eficiência.
"""
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
import pandas as pd
import openpyxl
from collections import Counter

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import automacao

def extrair_demanda_historica(caminho_arquivo, cad):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    registros = []
    
    # 1. Capital (PESO_CAP)
    if "PESO_CAP" in wb.sheetnames:
        ws = wb["PESO_CAP"]
        for r in range(5, ws.max_row + 1):
            loja_str = ws.cell(r, 2).value
            peso = ws.cell(r, 5).value
            rota_val = ws.cell(r, 6).value
            geracao = str(ws.cell(r, 8).value or "HOJE").strip().upper()
            
            if loja_str and peso is not None:
                f_cod = str(loja_str).split(" - ")[0].strip()
                if f_cod.isdigit() and len(f_cod) == 4 and not f_cod.startswith("14") and f_cod != "2061":
                    continue # ignorar lojas condominio
                
                try:
                    peso_val = float(peso)
                except (ValueError, TypeError):
                    continue
                
                info_cad = cad.get(f_cod, {})
                rota_padrao = info_cad.get("Rota Padrão") or automacao.normalizar_nome_rota(rota_val)
                regiao = info_cad.get("Região") or "CAPITAL"
                
                eh_hoje = ("HOJE" in geracao) or ("01" in geracao)
                status_dias = "01-DE HOJE" if eh_hoje else "02-NO PRAZO"
                
                registros.append({
                    "SECTOR": "CONGELADOS",
                    "FILIAL_PADRAO": f_cod,
                    "CLIENTE_PADRAO": str(loja_str),
                    "REGIAO_PADRAO": regiao.upper(),
                    "ROTA_PADRAO": automacao.normalizar_nome_rota(rota_padrao),
                    "PESOBRUTO": peso_val,
                    "DIAS_CARREGAMENTO": status_dias
                })

    # 2. Interior (PESO__INT)
    if "PESO__INT" in wb.sheetnames:
        ws = wb["PESO__INT"]
        for r in range(4, ws.max_row + 1):
            loja_str = ws.cell(r, 2).value
            peso = ws.cell(r, 5).value
            rota_val = ws.cell(r, 6).value
            geracao = str(ws.cell(r, 8).value or "HOJE").strip().upper()
            
            if loja_str and peso is not None:
                f_cod = str(loja_str).split(" - ")[0].strip()
                if f_cod.isdigit() and len(f_cod) == 4 and not f_cod.startswith("14") and f_cod != "2061":
                    continue
                    
                try:
                    peso_val = float(peso)
                except (ValueError, TypeError):
                    continue

                info_cad = cad.get(f_cod, {})
                rota_padrao = info_cad.get("Rota Padrão") or automacao.normalizar_nome_rota(rota_val)
                regiao = info_cad.get("Região") or "INTERIOR"
                
                eh_hoje = ("HOJE" in geracao) or ("01" in geracao)
                status_dias = "01-DE HOJE" if eh_hoje else "02-NO PRAZO"
                
                registros.append({
                    "SECTOR": "CONGELADOS",
                    "FILIAL_PADRAO": f_cod,
                    "CLIENTE_PADRAO": str(loja_str),
                    "REGIAO_PADRAO": regiao.upper(),
                    "ROTA_PADRAO": automacao.normalizar_nome_rota(rota_padrao),
                    "PESOBRUTO": peso_val,
                    "DIAS_CARREGAMENTO": status_dias
                })
                
    return pd.DataFrame(registros)

def extrair_rotas_manuais(caminho_arquivo):
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    rotas_manuais = {"CAPITAL": [], "INTERIOR": []}
    
    # Capital
    if "ROTAS_CAP" in wb.sheetnames:
        ws = wb["ROTAS_CAP"]
        for r in range(8, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(2, 11)]
            paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in vals[:4] if p is not None]
            veic = str(vals[6]).strip().upper().replace(".", "") if vals[6] else "N/D"
            if paradas:
                rotas_manuais["CAPITAL"].append({
                    "paradas": paradas,
                    "veiculo": veic
                })
                
    # Interior
    if "ROTA_INT" in wb.sheetnames:
        ws = wb["ROTA_INT"]
        for r in range(8, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(2, 11)]
            paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in vals[:4] if p is not None]
            veic = str(vals[6]).strip().upper().replace(".", "") if vals[6] else "N/D"
            if paradas:
                rotas_manuais["INTERIOR"].append({
                    "paradas": paradas,
                    "veiculo": veic
                })
                
    return rotas_manuais

def executar_benchmark():
    caminho_cad = automacao.ARQUIVO_ENDERECOS_PADRAO
    cad_df = pd.read_excel(caminho_cad)
    cad = {}
    for _, r in cad_df.iterrows():
        f = str(r["Filial"]).strip()
        cad[f] = {
            "Filial": f,
            "Nome Loja": str(r.get("Nome Loja", "")),
            "Região": str(r.get("Região", "")),
            "Rota Padrão": str(r.get("Rota Padrão", "")),
            "Restrição Veículo": str(r.get("Restrição Veículo", "Sem restrição")),
            "Bairro": str(r.get("Bairro", "")),
            "Cidade": str(r.get("Cidade", "")),
            "UF": str(r.get("UF", ""))
        }
        
    arquivos = [
        ("24/09/2026", r"C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 24-09.xlsx"),
        ("25/09/2026", r"C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 25-09 .xlsx"),
        ("26/09/2026", r"C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 26-09  - .xlsx")
    ]
    
    print("=" * 100)
    print("           BENCHMARK DE VALIDAÇÃO: ALGORITMO vs OPERAÇÃO REAL MANUAL (3 DIAS)           ")
    print("=" * 100)
    
    resumo_geral = []
    
    for dia_str, caminho in arquivos:
        if not os.path.exists(caminho):
            continue
            
        print(f"\n" + "#" * 90)
        print(f" >>> PROCESSANDO DATA: {dia_str} ({os.path.basename(caminho)})")
        print("#" * 90)
        
        df_demanda = extrair_demanda_historica(caminho, cad)
        rotas_manuais = extrair_rotas_manuais(caminho)
        
        # Executar alocação automática
        df_resumo_auto, df_det_auto, df_aguard_auto, df_sobras_auto = automacao.processar_alocacao_rotas(df_demanda, cad)
        
        p_total_dia = df_demanda["PESOBRUTO"].sum()
        p_auto = df_resumo_auto["Peso Total (Kg)"].sum() if not df_resumo_auto.empty else 0.0
        p_pend = df_aguard_auto["Peso Loja (Kg)"].sum() if not df_aguard_auto.empty else 0.0
        
        n_veics_auto = len(df_resumo_auto)
        n_veics_manual = len(rotas_manuais["CAPITAL"]) + len(rotas_manuais["INTERIOR"])
        
        print(f"\n[ESTATÍSTICAS DO DIA {dia_str}]:")
        print(f"   • Peso Bruto Total da Demanda:   {p_total_dia:>10,.2f} kg ({df_demanda['FILIAL_PADRAO'].nunique()} lojas)")
        print(f"   • Peso Roteirizado pelo Algoritmo: {p_auto:>10,.2f} kg ({p_auto/p_total_dia*100:>5.1f}%) em {n_veics_auto} veículos")
        print(f"   • Peso em Aguardo (Falta de Piso): {p_pend:>10,.2f} kg ({p_pend/p_total_dia*100:>5.1f}%)")
        print(f"   • Total de Veículos Manuais (Real): {n_veics_manual} veículos ({len(rotas_manuais['CAPITAL'])} Capital + {len(rotas_manuais['INTERIOR'])} Interior)")
        
        # Comparação da Frota
        frota_auto = Counter(df_resumo_auto["Veículo Ideal"].str.upper().str.replace(".", ""))
        frota_manual = Counter(v["veiculo"] for v in rotas_manuais["CAPITAL"] + rotas_manuais["INTERIOR"])
        
        print("\n   [DISTRIBUIÇÃO DA FROTA UTILIZADA]:")
        print(f"   {'Tipo de Veículo':<16} | {'Algoritmo':<12} | {'Operação Manual':<16} | {'Diferença':<10}")
        print("   " + "-" * 60)
        todos_tipos = sorted(list(set(list(frota_auto.keys()) + list(frota_manual.keys()))))
        for t in todos_tipos:
            qtd_a = frota_auto.get(t, 0)
            qtd_m = frota_manual.get(t, 0)
            dif = qtd_a - qtd_m
            print(f"   {t:<16} | {qtd_a:>8} un | {qtd_m:>12} un | {dif:>+8}")

        # Comparação por Região
        for reg in ["CAPITAL", "INTERIOR"]:
            sub_auto = df_resumo_auto[df_resumo_auto["Região"] == reg]
            sub_man = rotas_manuais[reg]
            
            print(f"\n   >>> DETALHE {reg}:")
            print(f"       Veículos Algoritmo: {len(sub_auto)} | Veículos Operador Manual: {len(sub_man)}")
            print("       Rotas montadas pelo Algoritmo:")
            for _, r in sub_auto.iterrows():
                paradas = [r[c] for c in ["1ª Entrega", "2ª Entrega", "3ª Entrega", "4ª Entrega"] if r[c]]
                print(f"          • {r['Rota de Execução']:<28} | {r['Veículo Ideal']:<8} | {r['Peso Total (Kg)']:>9.2f} kg | Lojas: {paradas}")
                
            print("       Rotas montadas pelo Operador Real:")
            for idx_m, r_m in enumerate(sub_man, 1):
                print(f"          • Carro {idx_m:02d} | {r_m['veiculo']:<8} | Lojas: {r_m['paradas']}")

        resumo_geral.append({
            "Data": dia_str,
            "Peso Total (Kg)": p_total_dia,
            "Peso Auto (Kg)": p_auto,
            "Veics Auto": n_veics_auto,
            "Veics Manual": n_veics_manual
        })

    print("\n" + "=" * 100)
    print("                              CONSOLIDADO DOS 3 DIAS TESTADOS                              ")
    print("=" * 100)
    df_consolidado = pd.DataFrame(resumo_geral)
    print(df_consolidado.to_string(index=False))

if __name__ == "__main__":
    executar_benchmark()
