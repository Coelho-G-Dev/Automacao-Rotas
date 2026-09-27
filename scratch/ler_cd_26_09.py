import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

path = r'C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 26-09  - .xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ws_cd = wb['CROSSDOCK']
print("=== CROSSDOCK 26/09 ===")
for r in range(1, ws_cd.max_row + 1):
    vals = [ws_cd.cell(r, c).value for c in range(1, ws_cd.max_column + 1)]
    if any(v is not None for v in vals):
        print(f"L{r:02d}: {[v for v in vals if v is not None]}")
