import openpyxl, sys
sys.stdout.reconfigure(encoding='utf-8')

path = r'C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 26-09  - .xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ws_pint = wb['PESO__INT']
print("=== PESO__INT 26/09 ===")
for r in range(5, ws_pint.max_row + 1):
    vals = [ws_pint.cell(r, c).value for c in range(1, 10)]
    if any(v is not None for v in vals):
        print(f"L{r:02d}: {vals[1:9]}")
