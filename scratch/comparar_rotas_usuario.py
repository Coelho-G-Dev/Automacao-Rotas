import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import openpyxl
import pandas as pd
from automacao import localizar_aba_e_cabecalho, carregar_dados_frios, obter_base_enderecos

# 1. Carregar arquivo do usuario
caminho_user = r"C:\Users\gabri\Desktop\Trabalho\Rotas\Rotas dia 24-09.xlsx"
wb_user = openpyxl.load_workbook(caminho_user, data_only=True)

# Lojas e pesos no PESO_CAP do usuario
ws_pcap = wb_user["PESO_CAP"]
lojas_user_cap = {}
for r in range(6, ws_pcap.max_row + 1):
    f_str = ws_pcap.cell(r, 2).value
    peso = ws_pcap.cell(r, 5).value
    rota_u = ws_pcap.cell(r, 6).value
    geracao = ws_pcap.cell(r, 8).value
    if f_str is not None and peso is not None:
        cod = str(f_str).split(" - ")[0].strip()
        lojas_user_cap[cod] = {
            "nome_completo": str(f_str),
            "peso": float(peso),
            "rota_user": str(rota_u),
            "geracao": str(geracao)
        }

# Rotas do usuario na Capital
ws_rcap = wb_user["ROTAS_CAP"]
rotas_user_cap = []
for r in range(8, ws_rcap.max_row + 1):
    p1 = ws_rcap.cell(r, 2).value
    p2 = ws_rcap.cell(r, 3).value
    p3 = ws_rcap.cell(r, 4).value
    p4 = ws_rcap.cell(r, 5).value
    veic = ws_rcap.cell(r, 9).value
    obs = ws_rcap.cell(r, 10).value
    paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in [p1, p2, p3, p4] if p is not None]
    if paradas:
        # Calcular peso somado das paradas
        peso_total_veic = 0.0
        for p in paradas:
            if p in lojas_user_cap:
                peso_total_veic += lojas_user_cap[p]["peso"]
            elif p == "96":
                # ver quanto foi
                peso_total_veic += lojas_user_cap.get(p, {}).get("peso", 0.0)
        rotas_user_cap.append({
            "veiculo": str(veic),
            "paradas": paradas,
            "peso_calculado": peso_total_veic,
            "obs": obs
        })

print("=== ROTAS MANUAIS DA CAPITAL FEITAS PELO USUÁRIO (24/09) ===")
for idx, r in enumerate(rotas_user_cap, 1):
    print(f"Carro {idx:02d} | Veículo: {r['veiculo']:<7} | Peso: {r['peso_calculado']:>10.2f} kg | Paradas: {r['paradas']}")

# Lojas no PESO__INT
ws_pint = wb_user["PESO__INT"]
lojas_user_int = {}
for r in range(6, ws_pint.max_row + 1):
    f_str = ws_pint.cell(r, 2).value
    peso = ws_pint.cell(r, 5).value
    rota_u = ws_pint.cell(r, 6).value
    geracao = ws_pint.cell(r, 8).value
    if f_str is not None and peso is not None:
        cod = str(f_str).split(" - ")[0].strip()
        lojas_user_int[cod] = {
            "nome_completo": str(f_str),
            "peso": float(peso),
            "rota_user": str(rota_u),
            "geracao": str(geracao)
        }

ws_rint = wb_user["ROTA_INT"]
rotas_user_int = []
for r in range(8, ws_rint.max_row + 1):
    p1 = ws_rint.cell(r, 2).value
    p2 = ws_rint.cell(r, 3).value
    p3 = ws_rint.cell(r, 4).value
    p4 = ws_rint.cell(r, 5).value
    veic = ws_rint.cell(r, 9).value
    paradas = [str(int(p)) if isinstance(p, (int, float)) else str(p).strip() for p in [p1, p2, p3, p4] if p is not None]
    if paradas:
        peso_total_veic = sum(lojas_user_int.get(p, {}).get("peso", 0.0) for p in paradas)
        rotas_user_int.append({
            "veiculo": str(veic),
            "paradas": paradas,
            "peso_calculado": peso_total_veic
        })

print("\n=== ROTAS MANUAIS DO INTERIOR FEITAS PELO USUÁRIO (24/09) ===")
for idx, r in enumerate(rotas_user_int, 1):
    print(f"Carro {idx:02d} | Veículo: {r['veiculo']:<7} | Peso: {r['peso_calculado']:>10.2f} kg | Paradas: {r['paradas']}")

ws_cd = wb_user["CROSSDOCK"]
rotas_user_cd = []
for r in range(5, ws_cd.max_row + 1):
    paradas = [str(int(ws_cd.cell(r, c).value)) if isinstance(ws_cd.cell(r, c).value, (int, float)) else str(ws_cd.cell(r, c).value).strip() for c in range(2, 9) if ws_cd.cell(r, c).value is not None]
    veic = ws_cd.cell(r, 10).value
    if paradas:
        rotas_user_cd.append({
            "veiculo": str(veic),
            "paradas": paradas
        })

print("\n=== ROTAS CROSSDOCK FEITAS PELO USUÁRIO (24/09) ===")
for idx, r in enumerate(rotas_user_cd, 1):
    print(f"Carro {idx:02d} | Veículo: {r['veiculo']:<7} | Paradas: {r['paradas']}")
